# -*- coding: utf-8 -*-
"""Genere les documents Word "Programme <joueur>" de l'ecole de tennis du Lambermont.

Le logo est place en haut du corps du document (et non dans la zone d'en-tete
de Word, qui s'affiche en grise tant qu'on edite le corps du document).
"""
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

LOGO = "/home/user/claude/assets/logo_lambermont.png"

from programmes_data import PROGRAMMES


def filet_bas(par, sz, couleur):
    """Filet sous le paragraphe, insere a la position imposee par le schema OOXML."""
    pPr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), couleur)
    bdr.append(bottom)
    style = pPr.find(qn("w:pStyle"))
    if style is not None:
        style.addnext(bdr)
    else:
        pPr.insert(0, bdr)


def generer(cle, data, dossier="/home/user/claude"):
    doc = Document()

    sec = doc.sections[0]
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(10)

    # En-tete dans le corps du document : logo + nom de l'ecole + filet
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_after = Pt(2)
    p_logo.add_run().add_picture(LOGO, width=Cm(8))

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(4)
    r = p_sub.add_run("École de tennis du Lambermont")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
    r.bold = True
    filet_bas(p_sub, 12, "55621E")

    # Titre
    titre = doc.add_paragraph()
    tr = titre.add_run(data["titre"])
    tr.bold = True
    tr.font.size = Pt(14)
    titre.paragraph_format.space_before = Pt(18)
    titre.paragraph_format.space_after = Pt(14)
    filet_bas(titre, 6, "000000")

    # Planning
    for ligne in data["planning"]:
        doc.add_paragraph(ligne)

    lignes_note = []
    if data.get("note"):
        lignes_note.append(data["note"])
    lignes_note += list(data.get("remarques", []))
    for i, texte in enumerate(lignes_note):
        note = doc.add_paragraph()
        nr = note.add_run(texte)
        nr.italic = True
        note.paragraph_format.space_before = Pt(6) if i == 0 else Pt(0)
        note.paragraph_format.space_after = Pt(14) if i == len(lignes_note) - 1 else Pt(4)
    if not lignes_note:
        doc.paragraphs[-1].paragraph_format.space_after = Pt(20)

    # Tableau
    lignes = [("Cours", "Prix normal (€)", "Prix adapté (€)")] + list(data["tableau"])
    table = doc.add_table(rows=len(lignes), cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    largeurs = [Cm(6.5), Cm(4.0), Cm(5.5)]
    for i, ligne in enumerate(lignes):
        cells = table.rows[i].cells
        for j, txt in enumerate(ligne):
            cells[j].width = largeurs[j]
            para = cells[j].paragraphs[0]
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after = Pt(2)
            run = para.add_run(txt)
            if i == 0 or i == len(lignes) - 1:
                run.bold = True

    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")

    chemin = "%s/Programme_%s.docx" % (dossier, cle)
    doc.save(chemin)
    return chemin


if __name__ == "__main__":
    cles = sys.argv[1:] or list(PROGRAMMES)
    for cle in cles:
        print(generer(cle, PROGRAMMES[cle]))
